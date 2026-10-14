#!/usr/bin/env python3
"""One-time seed of the Supplier Budgets tables from the tracker workbook.

    python3 supplier-budget/seed_from_workbook.py <Supplier_Budgets_Tracker.xlsx>

Writes (both git-ignored, supabase/data/ is in .gitignore):
    supabase/data/supplier_budget_seed.sql     paste into the Supabase SQL Editor AFTER the
                                               20261014210000_supplier_budget.sql migration
    supabase/data/supplier_budget_seed_report.txt  what was read, what was flagged

The repo is PUBLIC: the workbook holds dollars by vendor, so NEITHER the workbook nor the
generated SQL may ever be committed. The script never prints an amount to the terminal.

What it does
  Budget tab              -> budget_accounts (year 2026; the sheet's 0 budget = no budget set = NULL)
  Expenses 2026 tab       -> supplier_expenses, source 'import_2026', source_key 'Expenses 2026:<row>'
  Expenses 2025 tab       -> supplier_expenses, source 'import_2025' (rows DATED 2025 only: the 2026
                             rows on that tab are copies of the 2026 tab) linked to year-2025
                             account rows (inactive, no budget) so history stays apart from the budget
  Budget Breakdown tab    -> basis_text on the matching budget row (matched by supplier name; the
                             rest are listed, not guessed)
  An Expenses Account with no Budget row -> a placeholder budget_accounts row, needs_review = true,
  code + brand manager copied from its own expense rows, supplier NULL. Listed in the report,
  never mapped by guess. Gavin / Alexis fix them on the Budget tab of the site.
  Kohler Spend is ALWAYS total - billback in the database (trigger); the sheet's own cell is kept
  in sheet_values and every disagreement is listed in the report.
  Re-pasting the SQL is safe: existing source_keys / accounts are left alone (on conflict do nothing).
"""
import sys, json, re, datetime as dt
from collections import Counter, defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
OUT_SQL = ROOT / "supabase/data/supplier_budget_seed.sql"
OUT_REP = ROOT / "supabase/data/supplier_budget_seed_report.txt"

def key(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()

def txt(v):
    if v is None: return None
    s = str(v).strip()
    if s.endswith(".0") and re.fullmatch(r"\d+\.0", s): s = s[:-2]
    return s or None

def num(v):
    if v is None or v == "": return None
    if isinstance(v, (int, float)): return round(float(v), 2)
    s = str(v).replace("$", "").replace(",", "").strip()
    try: return round(float(s), 2)
    except ValueError: return None

def q(v):
    if v is None: return "null"
    if isinstance(v, bool): return "true" if v else "false"
    if isinstance(v, (int, float)): return repr(v)
    return "'" + str(v).replace("'", "''") + "'"

def brand_key(x):
    s = str(x or "").lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\b(company|co|inc|llc|ltd|brewing|brewery|brewers|beverage|beverages|group|usa|brands|imports?|exports?|projects|international|distributing|the|and)\b", " ", s)
    return re.sub(r"\s+", "", s)

def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rep = []
    say = rep.append

    # ---------------------------------------------------------------- Budget tab
    ws = wb["Budget"]
    accounts = {}          # (year, KEY) -> dict
    dup_accounts = []
    for i, r in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
        year, code, bm, sup, acct, vendor, desc, budget = r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8]
        if acct is None and code is None: continue
        if key(acct) in ("", "TOTAL", "TOTAL BUDGET"): continue
        try: year = int(year)
        except (TypeError, ValueError): year = 2026
        k = (year, key(acct))
        row = dict(year=year, code=txt(code), account=re.sub(r"\s+", " ", str(acct)).strip(), supplier=txt(sup),
                   brand_manager=txt(bm), description=txt(desc), vendor_default=txt(vendor),
                   budget_amount=num(budget) or None, needs_review=False, active=True, notes=None, basis=None, src=f"Budget:{i}")
        if row["budget_amount"] is not None and row["budget_amount"] <= 0: row["budget_amount"] = None
        if k in accounts:
            dup_accounts.append((i, row["account"], row["code"], accounts[k]["src"]))
            continue
        accounts[k] = row
    say(f"Budget tab: {len(accounts)} account rows (year 2026)")
    for d in dup_accounts: say(f"  DUPLICATE account on the Budget tab, row {d[0]} '{d[1]}' code {d[2]} -- kept {d[3]}, skipped this one")
    no_sup = [a for a in accounts.values() if not a["supplier"]]
    say(f"  {len(no_sup)} Budget rows have no Supplier (a brand manager reading by supplier will not see them until it is filled): " + "; ".join(f"{a['account']} ({a['code']})" for a in no_sup))
    no_budget = [a for a in accounts.values() if a["budget_amount"] is None]
    say(f"  {len(no_budget)} Budget rows have 0 / blank Budget Amount -> stored as no budget set: " + "; ".join(a["account"] for a in no_budget))
    codes = Counter(a["code"] for a in accounts.values() if a["code"])
    say("  codes shared by two accounts: " + ("; ".join(f"{c}: " + ", ".join(a["account"] for a in accounts.values() if a["code"] == c) for c, n in codes.items() if n > 1) or "none"))

    # ---------------------------------------------------------------- Budget Breakdown
    bb = wb["Budget Breakdown"] if "Budget Breakdown" in wb.sheetnames else None
    basis_rows = []
    if bb:
        for r in bb.iter_rows(min_row=3, values_only=True):
            brand, basis = r[0], r[1]
            if brand is None or key(brand) in ("BRAND",) or not basis: continue
            basis_rows.append((str(brand).strip(), str(basis).strip()))
        matched = 0; unmatched = []
        for brand, basis in basis_rows:
            bk = brand_key(brand)
            hits = [a for a in accounts.values() if a["supplier"] and (brand_key(a["supplier"]) == bk or (len(bk) >= 4 and (brand_key(a["supplier"]).startswith(bk) or bk.startswith(brand_key(a["supplier"])))))]
            if hits:
                for a in hits: a["basis"] = basis
                matched += 1
            else:
                unmatched.append(f"{brand} ({basis})")
        say(f"Budget Breakdown: {len(basis_rows)} basis lines, {matched} matched a Budget supplier; NOT matched (left for Alexis to type on the Budget tab): " + "; ".join(unmatched))

    # ---------------------------------------------------------------- Expenses
    def read_expenses(name, source, only_year=None):
        w = wb[name]
        out, flags = [], defaultdict(list)
        for i, r in enumerate(w.iter_rows(min_row=8, values_only=True), 8):
            r = list(r[:17]) + [None] * (17 - len(r[:17]))
            date, year, inv, vendor, desc, bm, items, total, spend, bbk, recv, ref, code, acct, fname, notes, _ = r
            if date is None and vendor is None and total is None and acct is None: continue
            if not isinstance(date, (dt.datetime, dt.date)):
                flags["no_date"].append(f"row {i}: Date '{date}' -- SKIPPED"); continue
            date = date.date() if isinstance(date, dt.datetime) else date
            if only_year and date.year != only_year:
                flags["other_year"].append(i); continue
            try: ycol = int(year) if year is not None else None
            except (TypeError, ValueError): ycol = None
            if ycol is not None and ycol != date.year:
                flags["year_col"].append(f"row {i}: Year column {ycol}, Date {date} -> year {date.year} used")
            total_n, spend_n, bb_n = num(total), num(spend), num(bbk)
            if bb_n is None:
                if bbk not in (None, ""): flags["billback_text"].append(f"row {i}: Billback Amt '{bbk}' read as 0")
                bb_n = 0.0
            derived = False
            if total_n is None:
                if spend_n is not None:
                    total_n = round(spend_n + bb_n, 2); derived = True
                    flags["total_derived"].append(f"row {i}: Total Amount blank -> Kohler Spend + Billback = total")
                else:
                    flags["total_missing"].append(f"row {i}: no Total Amount and no Kohler Spend -- SKIPPED ({txt(vendor)}, {date}, {txt(acct)})"); continue
            if spend_n is not None and abs(round(total_n - bb_n, 2) - spend_n) > 0.005 and not derived:
                flags["spend_mismatch"].append(f"row {i}: {date} {txt(vendor)} / {txt(acct)}: sheet Kohler Spend != Total - Billback (difference {round(spend_n - (total_n - bb_n), 2)}) -> computed value used")
            rv = key(recv)
            if rv in ("Y", "YES", "TRUE", "1"): received = True
            elif rv in ("N", "NO", "", "NONE", "0", "0.0", "FALSE"): received = False
            else: received = False; flags["received_odd"].append(f"row {i}: Billback Received? '{recv}' read as No")
            if rv not in ("Y", "N", ""): flags["received_odd2"].append(f"row {i}: '{recv}'")
            vend = txt(vendor) or "(No Vendor)"
            if vend == "(No Vendor)": flags["no_vendor"].append(i)
            acct_k = key(acct)
            if not acct_k:
                flags["no_account"].append(f"row {i}: no Account ({vend}, {date}, description '{txt(desc)}') -> filed under '(No Account)'")
                acct_k = "(NO ACCOUNT)"; acct_s = "(No Account)"
            else:
                acct_s = re.sub(r"\s+", " ", str(acct)).strip()
            sheet = {"Date": str(date), "Year": txt(year), "Invoice": txt(inv), "Vendor Name": txt(vendor), "Description": txt(desc),
                     "Brand Manager": txt(bm), "Items Purchased": txt(items), "Total Amount": num(total), "Kohler Spend": spend_n,
                     "Billback Amt": txt(bbk), "Billback Received?": txt(recv), "Reference #": txt(ref), "Code": txt(code),
                     "Account": txt(acct), "File Name": txt(fname), "Notes": txt(notes)}
            out.append(dict(src=f"{name}:{i}", source=source, date=date, year=date.year, invoice=txt(inv), vendor=vend,
                            items=txt(items), total=total_n, billback=bb_n, received=received, notes=txt(notes),
                            acct_k=acct_k, acct_s=acct_s, code=txt(code), bm=txt(bm), desc=txt(desc), fname=txt(fname), sheet=sheet))
        return out, flags

    e26, f26 = read_expenses("Expenses 2026", "import_2026")
    e25, f25 = read_expenses("Expenses 2025 (Original)", "import_2025", only_year=2025)
    say(f"Expenses 2026 tab: {len(e26)} rows read")
    say(f"Expenses 2025 (Original) tab: {len(e25)} rows dated 2025 read; {len(f25['other_year'])} rows dated 2026 on that tab SKIPPED (copies of the 2026 tab)")

    # placeholders for accounts missing from Budget
    missing = defaultdict(list)
    for e in e26 + e25:
        k = (e["year"], e["acct_k"])
        if k not in accounts:
            missing[k].append(e)
    for (year, ak), rows in sorted(missing.items()):
        codes = Counter(r["code"] for r in rows if r["code"]).most_common(1)
        bms = Counter(r["bm"] for r in rows if r["bm"]).most_common(1)
        names = Counter(r["acct_s"] for r in rows).most_common(1)
        accounts[(year, ak)] = dict(year=year, code=codes[0][0] if codes else None, account=names[0][0], supplier=None,
                                    brand_manager=bms[0][0] if bms else None, description=Counter(r["desc"] for r in rows if r["desc"]).most_common(1)[0][0] if any(r["desc"] for r in rows) else None,
                                    vendor_default=None, budget_amount=None, needs_review=(year == 2026), active=(year == 2026),
                                    notes=("Created by the seed: this Account appears on the Expenses tab but not on the Budget tab. Fill in Supplier / Budget or merge it." if year == 2026 else "2025 history (from the Expenses 2025 tab)"),
                                    basis=None, src=f"placeholder from {rows[0]['src']}")
    m26 = [(k, v) for k, v in missing.items() if k[0] == 2026]
    say(f"ACCOUNTS ON THE 2026 EXPENSES TAB WITH NO BUDGET ROW: {len(m26)} (placeholder rows, needs_review = true, supplier blank -- NOT mapped by guess):")
    for (year, ak), rows in sorted(m26):
        a = accounts[(year, ak)]
        say(f"  '{a['account']}'  code {a['code'] or '-'}  brand manager {a['brand_manager'] or '-'}  ({len(rows)} expense rows)")
    say(f"2025 history accounts created (inactive, no budget): {sum(1 for k in missing if k[0] == 2025)}")

    def flag_block(name, f):
        for k, label in [("year_col", "Year column disagrees with the Date"), ("total_derived", "Total Amount blank, derived from Kohler Spend + Billback"),
                         ("total_missing", "no amount at all, skipped"), ("spend_mismatch", "sheet Kohler Spend differs from Total - Billback (computed value used)"),
                         ("received_odd", "Billback Received? not Y/N"), ("billback_text", "Billback Amt was text"), ("no_account", "no Account"), ("no_date", "no Date")]:
            if f.get(k):
                say(f"{name}: {len(f[k])} rows -- {label}")
                for line in f[k]: say("    " + line)
        if f.get("no_vendor"): say(f"{name}: {len(f['no_vendor'])} rows with no Vendor Name -> '(No Vendor)' (rows {', '.join(map(str, f['no_vendor'][:40]))}{'...' if len(f['no_vendor']) > 40 else ''})")
    flag_block("Expenses 2026", f26)
    flag_block("Expenses 2025", f25)

    # Brand manager / code on an expense row vs the Budget row
    mism_bm, mism_code = [], []
    for e in e26:
        a = accounts[(e["year"], e["acct_k"])]
        if e["bm"] and a["brand_manager"] and key(e["bm"]) != key(a["brand_manager"]): mism_bm.append(f"{e['src']}: '{e['bm']}' on the row, Budget says '{a['brand_manager']}' ({a['account']})")
        if e["code"] and a["code"] and key(e["code"]) != key(a["code"]): mism_code.append(f"{e['src']}: code {e['code']} on the row, Budget says {a['code']} ({a['account']})")
    say(f"Expenses 2026 rows whose Brand Manager differs from the Budget row (the Budget row wins on the site): {len(mism_bm)}")
    for l in mism_bm: say("    " + l)
    say(f"Expenses 2026 rows whose Code differs from the Budget row (the Budget row wins on the site): {len(mism_code)}")
    for l in mism_code[:60]: say("    " + l)
    if len(mism_code) > 60: say(f"    ... and {len(mism_code) - 60} more")

    vendors = Counter(e["vendor"] for e in e26 + e25)
    say(f"Vendors: {len(vendors)} distinct names (filled by the trigger as rows are inserted)")

    # ---------------------------------------------------------------- SQL
    L = ["-- Supplier Budgets seed, generated " + dt.datetime.now().strftime("%Y-%m-%d %H:%M") + " by supplier-budget/seed_from_workbook.py",
         "-- NEVER COMMIT THIS FILE (supabase/data/ is git-ignored): it holds dollars by vendor.",
         "-- Run once in the Supabase SQL Editor after 20261014210000_supplier_budget.sql. Safe to run again: existing rows are left alone.",
         "begin;", "select set_config('kdh.seeding', '1', false);"]
    for a in sorted(accounts.values(), key=lambda a: (a["year"], a["account"])):
        L.append("insert into public.budget_accounts (year, code, account, supplier, brand_manager, description, vendor_default, budget_amount, basis_text, active, needs_review, notes, created_by, updated_by) values (" +
                 ", ".join(q(x) for x in [a["year"], a["code"], a["account"], a["supplier"], a["brand_manager"], a["description"], a["vendor_default"], a["budget_amount"], a["basis"], a["active"], a["needs_review"], a["notes"], "import", "import"]) +
                 ") on conflict (year, account) do nothing;")
    for e in e26 + e25:
        a = accounts[(e["year"], e["acct_k"])]
        L.append("insert into public.supplier_expenses (account_id, expense_date, year, invoice, vendor, items_purchased, total_amount, billback_amount, billback_received, received_date, notes, source, source_key, sheet_values, created_by, updated_by) select id, " +
                 ", ".join(q(x) for x in [str(e["date"]), e["year"], e["invoice"], e["vendor"], e["items"], e["total"], e["billback"], e["received"], None, e["notes"], e["source"], e["src"], json.dumps(e["sheet"], ensure_ascii=False), "import", "import"]) +
                 f" from public.budget_accounts where year = {a['year']} and account = {q(a['account'])} on conflict (source_key) do nothing;")
    L.append("commit;")
    OUT_SQL.write_text("\n".join(L) + "\n")
    say(f"SQL: {len(accounts)} account rows + {len(e26) + len(e25)} expense rows -> {OUT_SQL.relative_to(ROOT)}")
    OUT_REP.write_text("\n".join(rep) + "\n")
    print("\n".join(l for l in rep if not l.startswith("    ")))
    print(f"full report (with the row-by-row lists): {OUT_REP.relative_to(ROOT)}")

if __name__ == "__main__":
    if len(sys.argv) != 2: sys.exit(__doc__)
    main(sys.argv[1])
