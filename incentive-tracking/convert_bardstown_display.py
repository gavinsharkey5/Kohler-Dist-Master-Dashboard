"""Bardstown / Green River Display & Activation (Sept - Oct 2026): merge iSellBeer
reports into the two files build_bardstown_display() reads.

    python3 incentive-tracking/convert_bardstown_display.py Promos_Report_NN.xlsx Report_NN.xlsx

Either report, in any order, as many as you like. The kind is read from the header:
  * Displays report ("Brand Family", "SKU", "Quantity" columns)  -> data/bardstown_displays.csv
    (off-premise stack displays; Quantity = cases on display)
  * Promos report ("Promotion type", "Elements" columns)         -> data/bardstown_activations.csv
    (on-premise activations / feature drinks and their documentation)

Only Bardstown Bourbon Company rows are kept. Rows MERGE onto what is on file: the key is
the photo link + brand (iSellBeer's Promo # / row # renumbers on every pull, so it is never
the key), a row seen again replaces the old copy, and nothing already on file is dropped --
a weekly pull of just the newest days is safe. Photo links are the hyperlink TARGETS.
"""
import csv
import sys
from pathlib import Path

import openpyxl

DATA = Path(__file__).parent / "data"
SUPPLIER = "BARDSTOWN BOURBON COMPANY"
OUT = {
    "displays": (DATA / "bardstown_displays.csv",
                 ["date", "account", "dba", "city", "photo_taker", "role", "brand_family", "brand", "sku", "quantity", "photo"]),
    "promos": (DATA / "bardstown_activations.csv",
               ["date", "account", "dba", "city", "photo_taker", "role", "promotion_type", "theme", "elements", "brand", "quantity", "photo"]),
}


def read(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    head = [str(c.value or "").strip() for c in ws[1]]
    col = {h.lower(): i for i, h in enumerate(head)}
    if "brand family" in col and "sku" in col:
        kind = "displays"
    elif "promotion type" in col:
        kind = "promos"
    else:
        raise SystemExit(f"{path}: not an iSellBeer Displays or Promos report (header: {head})")
    rows = []
    for r in ws.iter_rows(min_row=2):
        g = lambda name: r[col[name]].value if name in col else None
        if not any(c.value not in (None, "") for c in r):
            continue
        if str(g("supplier") or "").strip().upper() != SUPPLIER:
            continue
        pc = r[col["photo"]]
        photo = pc.hyperlink.target if pc.hyperlink else str(pc.value or "")
        base = {"date": str(g("date/time") or "").strip(), "account": str(g("account #") or "").strip(),
                "dba": str(g("dba") or "").strip(), "city": str(g("city") or "").strip(),
                "photo_taker": str(g("photo taker") or "").strip(),
                "role": str(g("photo taker's role") or "").strip(), "brand": str(g("brand") or "").strip(),
                "quantity": "" if g("quantity") in (None, "") else str(g("quantity")), "photo": photo}
        if kind == "displays":
            base.update(brand_family=str(g("brand family") or "").strip(), sku=str(g("sku") or "").strip())
        else:
            base.update(promotion_type=str(g("promotion type") or "").strip(), theme=str(g("theme") or "").strip(),
                        elements=str(g("elements") or "").strip())
        rows.append(base)
    return kind, rows


def main(paths):
    if not paths:
        raise SystemExit(__doc__)
    for p in paths:
        kind, rows = read(p)
        out, cols = OUT[kind]
        have = {}
        if out.exists():
            with open(out, newline="", encoding="utf-8") as f:
                for r in csv.DictReader(f):
                    have[(r["photo"], r["brand"])] = r
        before = len(have)
        for r in rows:
            have[(r["photo"], r["brand"])] = r
        with open(out, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            for r in sorted(have.values(), key=lambda r: (r["date"][6:10] + r["date"][:5] + r["date"][10:], r["account"], r["brand"])):
                w.writerow({c: r.get(c, "") for c in cols})
        print(f"{Path(p).name}: {kind} report, {len(rows)} Bardstown rows -> {out.name} "
              f"({before} on file, {len(have)} now)")


if __name__ == "__main__":
    main(sys.argv[1:])
