"""Write shared/data/tap-rules.json -- the US vs THEM rulebook for tap lines a
rep captures in the Hub (Account -> Add Photos -> Tap Handles).

Source: iSellBeer_TAPS_US_THEM_Mediator.xlsx, the same workbook the Tap
Tracker's audit replays (audit_engine.py). Nothing is re-derived: each picker
brand is resolved to its Encompass brand family with audit_engine's own steps
(Brand Family Territory -> Brand Crosswalk -> Brands (Enc)), and the label for
a family in an area is the workbook's "Master - US vs THEM" Final
Determination. The page then mirrors audit():
  brand contains "(IN-HOUSE)"           -> US
  family "No Encompass Match"           -> THEM (Kohler does not carry it)
  family "Not Mapped" / no area / no row -> not determined: the rep says
  (families are compared in upper case, as audit_engine does)
  otherwise                              -> master[family|AREA]
AREA = the account's distribution area (Bergen, Passaic, Passaic-FF, Essex,
Hudson, Union, Sussex, Morris 1/2/3); an account filed under "Sales" uses its
county, except Morris / Default (no single area -> not determined), exactly as
audit_engine column U does.

Run after the mediator workbook changes:  python3 isellbeer/tap-survey-tracking/build_tap_rules.py
"""
import json, warnings, datetime
from pathlib import Path
import openpyxl
from audit_engine import build, s

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
SRC = HERE / "iSellBeer_TAPS_US_THEM_Mediator.xlsx"
OUT = ROOT / "shared" / "data" / "tap-rules.json"

def family_for(K, L, refs):
    """audit_engine.audit() columns P..S for a brand K / brand family L."""
    terr, cw, brands, cust, master = refs
    P = L if L.upper() in terr else ""
    lookup = (L if L else K).upper()
    Q, R = cw.get(lookup, ("Not in Crosswalk", ""))
    if P: return P
    if Q == "Mapped": return R
    if Q == "No Encompass Match": return "No Encompass Match"
    return brands.get(lookup, "Not Mapped")

def main():
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    refs = build(wb)
    # families compare in UPPER CASE, as audit_engine's master lookup does (V.upper())
    master, name = {}, {}
    for r in wb["Master - US vs THEM"].iter_rows(min_row=2, values_only=True):
        fam, area, det = s(r[0]), s(r[1]).upper(), s(r[5])
        if fam and area and det in ("US", "THEM"):
            name.setdefault(fam.upper(), fam)
            master.setdefault(fam.upper(), {}).setdefault(area, det)     # first hit, as audit_engine
    picks = {}
    # the brands reps saw on taps in iSellBeer, resolved through the crosswalk
    for r in wb["Brand Segments (iSell)"].iter_rows(min_row=2, values_only=True):
        K, L = s(r[1]), s(r[2])
        if K and K.upper() not in picks: picks[K.upper()] = [K, family_for(K, L, refs).upper()]
    # every brand actually surveyed on a tap (the survey sheet), same resolution
    for r in wb["iSellBeer Import Template"].iter_rows(min_row=2, values_only=True):
        K, L = s(r[10]), s(r[11])
        if K and K.upper() not in picks: picks[K.upper()] = [K, family_for(K, L, refs).upper()]
    # Encompass's own brands (what Kohler carries), family as Encompass files it
    for r in wb["Brands (Enc)"].iter_rows(min_row=2, values_only=True):
        B, F, st = s(r[2]), s(r[3]), s(r[5])
        if B and F and st.lower() != "inactive" and B.upper() not in picks and B.lower() != "misc":
            picks[B.upper()] = [B, F.upper()]
    keys = sorted({f for _, f in picks.values()} | set(master))
    idx = {f: i for i, f in enumerate(keys)}
    fams = [name.get(k, k.title() if k not in ("NO ENCOMPASS MATCH", "NOT MAPPED") else k.title()) for k in keys]
    out = {
        "source": SRC.name,
        "built": datetime.date.today().isoformat(),
        "note": "US vs THEM for Hub tap lines; same rulebook and steps as the Tap Tracker audit (audit_engine.py).",
        "families": fams,
        "master": {str(idx[f]): v for f, v in master.items()},
        "brands": sorted(([b, idx[f]] for b, f in picks.values()), key=lambda x: x[0].upper()),
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, separators=(",", ":")))
    det = sum(1 for _, f in picks.values() if f in master)
    print(f"tap-rules.json: {len(picks)} brands ({det} with a territory ruling), {len(master)} families x areas, {OUT.stat().st_size // 1024} KB")

if __name__ == "__main__":
    main()
