#!/usr/bin/env python3
"""Builds Summer_Display_Auction_Billback_Jul-Aug_2026.xlsx -- the supplier
billback sheet for the July/August 2026 Summer Display Auction -- from the
recap workbook make_display_recap.py already writes.

    python3 make_display_billback.py

It reads the recap's 'Display Detail' tab (one row per photo submission,
already tiered, classed Priority / All Other and credited to one supplier by
the tracker's own scoring), so the billback cannot disagree with the recap or
the leaderboard about what a display is. The only thing added here is money:

    Gold   (70+ cases)   Priority $50   All Other $35
    Silver (40-69)       Priority $35   All Other $25
    Bronze (20-39)       Priority $20   All Other $15
    Other  (10-19)       any brand $10
    under 10 cases       $0 (does not qualify)

Every count and dollar on the summary tabs is a COUNTIFS / SUMIFS over the
detail tab, so filtering or excluding a row there flows through. The 'Bill?'
column on the detail starts at "Yes" for every qualifying display; set it to
"No" (e.g. for a confirmed duplicate) and the summaries drop it.

The output carries dollar amounts and lives in a PUBLIC repo -- it is listed
in .gitignore; send it, don't commit it.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
SRC = HERE / "Summer_Display_Auction_Jul-Aug_2026.xlsx"
OUT = HERE / "Summer_Display_Auction_Billback_Jul-Aug_2026.xlsx"

TIERS = [(4, "Gold", "70+ cases"), (3, "Silver", "40-69 cases"),
         (2, "Bronze", "20-39 cases"), (1, "All Other Displays", "10-19 cases, any brand")]
RATE = {(4, "Priority"): 50, (4, "All Other"): 35,
        (3, "Priority"): 35, (3, "All Other"): 25,
        (2, "Priority"): 20, (2, "All Other"): 15,
        (1, "Priority"): 10, (1, "All Other"): 10}

HEAD = PatternFill("solid", fgColor="0B2466")
SUB = PatternFill("solid", fgColor="E8EEF8")
TOT = PatternFill("solid", fgColor="FFF4E0")
GREY = Font(color="808080")
WHITE_B = Font(bold=True, color="FFFFFF")
B = Font(bold=True)
MONEY = '"$"#,##0'
THIN = Border(bottom=Side(style="thin", color="C8CED8"))
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")


def head_row(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row, i, t)
        c.fill, c.font, c.alignment = HEAD, WHITE_B, WRAP
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="0B2466")
    ws["A2"] = sub
    ws["A2"].font = GREY


def main():
    src = openpyxl.load_workbook(SRC, read_only=True)
    rows = list(src["Display Detail"].iter_rows(values_only=True))
    h = rows[0]
    D = [dict(zip(h, r)) for r in rows[1:] if r[0]]

    wb = openpyxl.Workbook()
    summ = wb.active
    summ.title = "Billback Summary"
    sup = wb.create_sheet("Billback by Supplier")
    per = wb.create_sheet("By Rep or Associate")
    det = wb.create_sheet("Display Detail")
    notes = wb.create_sheet("Notes")

    # ---------------- detail ----------------
    cols = ["Display ID", "Date", "Rep / Associate", "Role", "District Manager",
            "Account #", "Account (DBA)", "City", "Supplier Group", "Lead Brand",
            "All Brands in Display", "Cases", "Display Tier", "Tier #",
            "Priority / All Other", "Rate", "Bill?", "Payout", "Duplicate Flag", "Photo"]
    head_row(det, 1, cols, [9, 11, 18, 13, 16, 10, 30, 14, 20, 30, 40, 8, 15, 7,
                            13, 8, 7, 10, 18, 11])
    det.freeze_panes = "A2"
    photo_links = {}
    # the recap's photo cells carry the hyperlink; read them with links
    full = openpyxl.load_workbook(SRC)["Display Detail"]
    pcol = h.index("Photo Link") + 1
    for r in range(2, full.max_row + 1):
        c = full.cell(r, pcol)
        if c.hyperlink:
            photo_links[full.cell(r, 1).value] = c.hyperlink.target

    D.sort(key=lambda d: (-(d["Tier #"] or 0), d["Supplier Group"], d["Date / Time"]))
    for i, d in enumerate(D, 2):
        tier = d["Tier #"] or 0
        cls = d["Point Class"]
        rate = RATE.get((tier, cls), 0)
        dup = d["Duplicate Flag"] if d["Duplicate Flag"] == "Possible duplicate" else ""
        vals = [d["Display ID"], d["Date"], d["Rep / Associate"], d["Role"],
                d["District Manager"], d["Account #"], d["Account (DBA)"], d["City"],
                d["Supplier Group"], d["Lead Brand"], d["All Brands in Display"],
                d["Cases"], d["Display Tier"], tier, cls if tier else "",
                rate, "Yes" if tier else "No", f'=IF(Q{i}="Yes",P{i},0)', dup]
        for j, v in enumerate(vals, 1):
            det.cell(i, j, v)
        det.cell(i, 16).number_format = MONEY
        det.cell(i, 18).number_format = MONEY
        if d["Display ID"] in photo_links:
            c = det.cell(i, 20, "View Photo")
            c.hyperlink = photo_links[d["Display ID"]]
            c.font = Font(color="2866C0", underline="single")
        if not tier:
            for j in range(1, 20):
                det.cell(i, j).font = GREY
    last = len(D) + 1
    det.auto_filter.ref = f"A1:T{last}"
    rng = lambda col: f"'Display Detail'!${col}$2:${col}${last}"
    SUPC, TIERC, CLSC, BILL, PAY, ROLE, WHO = (rng("I"), rng("N"), rng("O"),
                                               rng("Q"), rng("R"), rng("D"), rng("C"))

    # ---------------- summary: tier x class grid ----------------
    title(summ, "Summer Display Auction -- Supplier Billback",
          "Kohler iSell Beer Summer Display Auction - 07/01/2026 to 08/31/2026 - "
          "one row per photo submission, counted once in its case-size tier")
    head_row(summ, 4, ["Display Tier", "Cases", "Priority / All Other", "Displays",
                       "Rate", "Payout"], [22, 22, 20, 12, 10, 14])
    r = 5
    first = r
    for tier, name, rule in TIERS:
        classes = ["Priority", "All Other"]
        for cls in classes:
            label = cls if tier > 1 else ("Priority brand" if cls == "Priority" else "All other brand")
            summ.cell(r, 1, name)
            summ.cell(r, 2, rule)
            summ.cell(r, 3, label)
            summ.cell(r, 4, f'=COUNTIFS({TIERC},{tier},{CLSC},"{cls}",{BILL},"Yes")')
            summ.cell(r, 5, RATE[(tier, cls)]).number_format = MONEY
            summ.cell(r, 6, f"=D{r}*E{r}").number_format = MONEY
            for j in range(1, 7):
                summ.cell(r, j).border = THIN
            r += 1
    summ.cell(r, 1, "Total").font = B
    summ.cell(r, 4, f"=SUM(D{first}:D{r-1})").font = B
    c = summ.cell(r, 6, f"=SUM(F{first}:F{r-1})")
    c.font, c.number_format = B, MONEY
    for j in range(1, 7):
        summ.cell(r, j).fill = TOT
    tot_row = r
    r += 2
    summ.cell(r, 1, "Excluded from billing").font = B
    summ.cell(r + 1, 1, "Under 10 cases (does not qualify)")
    summ.cell(r + 1, 4, f"=COUNTIF({TIERC},0)")
    summ.cell(r + 2, 1, 'Qualifying but marked Bill? = "No"')
    summ.cell(r + 2, 4, f'=COUNTIFS({TIERC},">0",{BILL},"No")')
    summ.cell(r + 3, 1, "Flagged possible duplicates (still billed)")
    summ.cell(r + 3, 4, f'=COUNTIFS(\'Display Detail\'!$S$2:$S${last},"Possible duplicate",{BILL},"Yes")')
    summ.cell(r + 3, 6, f'=SUMIFS({PAY},\'Display Detail\'!$S$2:$S${last},"Possible duplicate")').number_format = MONEY
    r += 5
    summ.cell(r, 1, "By role").font = B
    for k, role in enumerate(["Sales Rep", "Sales Associate"], 1):
        summ.cell(r + k, 1, role)
        summ.cell(r + k, 4, f'=COUNTIFS({ROLE},"{role}",{BILL},"Yes",{TIERC},">0")')
        summ.cell(r + k, 6, f'=SUMIFS({PAY},{ROLE},"{role}")').number_format = MONEY
    summ.freeze_panes = "A5"

    # ---------------- by supplier ----------------
    title(sup, "Billback by Supplier",
          "Each display is billed to the one supplier holding the most cases in the photo. "
          "Counts and dollars come from the Display Detail tab.")
    labels = ["Supplier Group"]
    for tier, name, _ in TIERS[:3]:
        labels += [f"{name}\nPriority", f"{name}\nAll Other"]
    labels += ["All Other Displays\n10-19 cases", "Total\nDisplays", "Total\nPayout"]
    head_row(sup, 4, labels, [34] + [11] * 7 + [11, 13])
    suppliers = sorted({d["Supplier Group"] for d in D if d["Tier #"]},
                       key=lambda s: (s.startswith("All Other"), s))
    r = 5
    for s in suppliers:
        sup.cell(r, 1, s)
        col = 2
        pay_terms = []
        for tier, _, _ in TIERS[:3]:
            for cls in ("Priority", "All Other"):
                sup.cell(r, col, f'=COUNTIFS({SUPC},$A{r},{TIERC},{tier},{CLSC},"{cls}",{BILL},"Yes")')
                pay_terms.append(f"{get_column_letter(col)}{r}*{RATE[(tier, cls)]}")
                col += 1
        sup.cell(r, col, f'=COUNTIFS({SUPC},$A{r},{TIERC},1,{BILL},"Yes")')
        pay_terms.append(f"{get_column_letter(col)}{r}*10")
        sup.cell(r, col + 1, f"=SUM(B{r}:{get_column_letter(col)}{r})")
        c = sup.cell(r, col + 2, "=" + "+".join(pay_terms))
        c.number_format = MONEY
        for j in range(1, col + 3):
            sup.cell(r, j).border = THIN
        r += 1
    sup.cell(r, 1, "Total").font = B
    for j in range(2, 11):
        c = sup.cell(r, j, f"=SUM({get_column_letter(j)}5:{get_column_letter(j)}{r-1})")
        c.font = B
        if j == 10:
            c.number_format = MONEY
    for j in range(1, 11):
        sup.cell(r, j).fill = TOT
    sup.cell(r + 1, 1, "Check: equals the Billback Summary total").font = GREY
    sup.cell(r + 1, 10, f"=J{r}-'Billback Summary'!F{tot_row}").number_format = MONEY
    sup.freeze_panes = "B5"

    # ---------------- by person ----------------
    title(per, "Payout by Rep or Associate", "Qualifying displays billed, by who took the photo.")
    head_row(per, 4, ["Rep / Associate", "Role", "District Manager", "Gold", "Silver",
                      "Bronze", "All Other Displays", "Displays", "Payout"],
             [22, 15, 18, 9, 9, 9, 12, 10, 12])
    # one row per name (a person can appear under two DMs; SUMIFS is by name)
    seen = {}
    for d in D:
        if d["Tier #"]:
            seen.setdefault(d["Rep / Associate"], (d["Role"], set()))[1].add(d["District Manager"] or "")
    people = sorted(((n, rl, ", ".join(sorted(dms))) for n, (rl, dms) in seen.items()),
                    key=lambda p: (p[1], p[0]))
    r = 5
    for name, role, dm in people:
        per.cell(r, 1, name)
        per.cell(r, 2, role)
        per.cell(r, 3, dm)
        for k, tier in enumerate([4, 3, 2, 1]):
            per.cell(r, 4 + k, f'=COUNTIFS({WHO},$A{r},{TIERC},{tier},{BILL},"Yes")')
        per.cell(r, 8, f"=SUM(D{r}:G{r})")
        per.cell(r, 9, f"=SUMIFS({PAY},{WHO},$A{r})").number_format = MONEY
        r += 1
    per.cell(r, 1, "Total").font = B
    for j in range(4, 10):
        c = per.cell(r, j, f"=SUM({get_column_letter(j)}5:{get_column_letter(j)}{r-1})")
        c.font = B
        if j == 9:
            c.number_format = MONEY
    for j in range(1, 10):
        per.cell(r, j).fill = TOT
    per.freeze_panes = "A5"

    # ---------------- notes ----------------
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 110
    lines = [
        ("Source", "Summer_Display_Auction_Jul-Aug_2026.xlsx (Display Detail tab), built by "
                   "make_display_recap.py from the full-period iSellBeer export. 690 submissions, "
                   "510 qualifying -- ties to the live leaderboard."),
        ("Window", "07/01/2026 - 08/31/2026."),
        ("One display", "One photo submission (photo taker + account + date/time). All SKUs in the photo "
                        "are summed; the tier is set by that total."),
        ("Payout table", "Gold 70+: Priority $50 / All Other $35. Silver 40-69: $35 / $25. "
                         "Bronze 20-39: $20 / $15. 10-19 cases, any brand: $10. Under 10: not paid."),
        ("Priority brands", "Constellation, Molson Coors, HUSA, Boston Beer, Mark Anthony, Yuengling, "
                            "Garage, Sapporo, Pabst, Montauk, Carbliss (program slide), classed per brand "
                            "by the tracker. A display mixing both sides takes the side with more cases."),
        ("Supplier billed", "The supplier group with the most cases in the photo. Displays with more than "
                            "one supplier are listed on the recap's Multi-Supplier Audit tab."),
        ("Possible duplicates", "67 qualifying displays are flagged in the recap (same account and day, "
                                "often a rep and an associate on the same display). They are BILLED here "
                                "unless you set Bill? = No on the Display Detail tab."),
        ("Bill? column", "Set to No to drop a display from every total; the summaries recalculate."),
    ]
    notes["A1"] = "Notes"
    notes["A1"].font = Font(bold=True, size=14, color="0B2466")
    for i, (a, b) in enumerate(lines, 3):
        notes.cell(i, 1, a).font = B
        notes.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)

    # printed check, computed in Python (independent of the formulas)
    total = sum(RATE.get(((d["Tier #"] or 0), d["Point Class"]), 0) for d in D)
    q = sum(1 for d in D if d["Tier #"])
    print(f"{OUT.name}: {len(D)} submissions, {q} qualifying, ${total:,} billback")


if __name__ == "__main__":
    main()
