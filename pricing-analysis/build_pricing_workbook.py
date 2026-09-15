#!/usr/bin/env python3
"""Build the Pricing / Deal Level workbook from the Encompass RDE export
"Pricing Analysis RDE by Month".

    python3 build_pricing_workbook.py [data/Pricing_Analysis_RDE_by_Month.csv] [out.xlsx]

One export row = one product x one deal level (Min Purchase Qty) x one deal
window, with the cases sold at that level. The workbook keeps the export's
Gross $ and Margin % as the source of truth (they are correct even where the
FOB / laid-in are quoted per bottle, e.g. spirits) and lays the levels out as
a price ladder per product, plus INFERRED mix-and-match families -- the
export carries no deal-group field, so families are derived from shared
ladders of 20+ case levels within a supplier. Confirm against the supplier
deal sheets before treating a family as fact.
"""
import re, sys, os
from collections import defaultdict, Counter
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table  # noqa (not used; autofilter instead)

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'data', 'Pricing_Analysis_RDE_by_Month.csv')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, 'Pricing_Deal_Levels_Aug_2026.xlsx')

FONT = 'Arial'
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color='FFFFFF')
F_TITLE = Font(name=FONT, size=14, bold=True)
F_NOTE = Font(name=FONT, size=10, italic=True, color='555555')
FILL_HDR = PatternFill('solid', fgColor='1F3864')
FILL_FAM = PatternFill('solid', fgColor='D9E1F2')
FILL_TOT = PatternFill('solid', fgColor='EDEDED')
FILL_INPUT = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

FMT_PRICE = '$#,##0.00;($#,##0.00);-'
FMT_MONEY = '$#,##0;($#,##0);-'
FMT_PCT = '0.0%;(0.0%);-'
FMT_INT = '#,##0;(#,##0);-'
FMT_DATE = 'm/d/yyyy'

# ----------------------------------------------------------------------------
# 1. Load + clean
# ----------------------------------------------------------------------------
def money(s):
    s = s.astype(str).str.replace('[$,]', '', regex=True).str.replace('(', '-', regex=False).str.replace(')', '', regex=False)
    return pd.to_numeric(s, errors='coerce')

raw = pd.read_csv(SRC)
df = raw.copy()
for c in ['Case Unit Price', 'FOB', 'Tax', 'Participation', 'Laid-in Cost', 'Gross']:
    df[c] = money(df[c])
df['Margin %'] = pd.to_numeric(df['Margin %'].astype(str).str.rstrip('%'), errors='coerce') / 100
df['Deal Start Date'] = pd.to_datetime(df['Deal Start Date'])
df['Deal End Date'] = pd.to_datetime(df['Deal End Date'])
df = df.rename(columns={'Min Purchase Qty': 'lvl', 'Case Unit Price': 'price', 'Deal Start Date': 'start', 'Deal End Date': 'end',
                        'Laid-in Cost': 'laidin', 'Margin %': 'margin'})
df['pnum'] = df['Product Num & Name'].str.extract(r'^(\d+)\s')[0].astype(int)
df['pname'] = df['Product Num & Name'].str.replace(r'^\d+\s', '', regex=True).str.strip()

def split_pkg(n):
    m = re.search(r'\b\d[\d./]*\s?(oz|mL|ML|L|Gal|BBL)\b.*$', n)
    if m and m.start() > 0:
        return n[:m.start()].strip(), m.group(0).strip()
    m = re.search(r'\bKeg\b.*$', n)
    if m and m.start() > 0:
        return n[:m.start()].strip(), m.group(0).strip()
    return n, ''
df[['brand', 'pkg']] = df['pname'].apply(lambda n: pd.Series(split_pkg(n)))
df['rev'] = df['Cases'] * df['price']
df['laid_calc'] = df['FOB'] + df['Tax'] - df['Participation']

# ---- flags ------------------------------------------------------------------
dupmask = df.duplicated(['pnum', 'lvl', 'start', 'end'], keep=False)
def flags(r, dup):
    f = []
    if r['Cases'] < 0: f.append('Negative cases (credit / return)')
    if pd.isna(r['FOB']): f.append('No cost on file')
    if r['margin'] < 0: f.append('Negative margin')
    if pd.notna(r['laidin']) and abs(r['laid_calc'] - r['laidin']) > 0.02:
        f.append('Laid-in != FOB+Tax-Part')
    if dup: f.append('Duplicate level/window (price override?)')
    return '; '.join(f)
df['flags'] = [flags(r, d) for r, d in zip(df.to_dict('records'), dupmask)]

# ---- inferred mix & match families -----------------------------------------
MM_MIN = 20  # levels at/above this are treated as program (mix-and-match) rungs
sig = df[df['lvl'] >= MM_MIN].groupby(['Supplier', 'pnum'])['lvl'].apply(lambda s: frozenset(s)).to_dict()
canon = defaultdict(list)
for (sup, p), s in sig.items():
    canon[(sup, s)].append(p)
canonical = {k: v for k, v in canon.items() if len(v) >= 2}

def jacc(a, b):
    return len(a & b) / len(a | b)

family_of = {}
for (sup, p), s in sig.items():
    if (sup, s) in canonical:
        family_of[p] = (sup, s); continue
    best, bj = None, 0.0
    for (csup, cs) in canonical:
        if csup != sup: continue
        j = jacc(s, cs)
        if j > bj or (j == bj and best is not None and len(cs) > len(best[1])):
            best, bj = (csup, cs), j
    family_of[p] = best if best is not None and bj >= 1 / 3 else (sup, s)

def fam_name(sup, s):
    return f"{sup} | {'/'.join(str(t) for t in sorted(s))}"

df['family'] = [fam_name(*family_of[p]) if p in family_of else f"{sup} | no {MM_MIN}+ case level" for p, sup in zip(df['pnum'], df['Supplier'])]

# ---- sort: supplier by revenue, brand by revenue, product by cases, level asc
sup_rank = df.groupby('Supplier')['rev'].sum().rank(ascending=False, method='first')
brand_rank = df.groupby(['Supplier', 'brand'])['rev'].sum().rank(ascending=False, method='first')
prod_rank = df.groupby('pnum')['Cases'].sum().rank(ascending=False, method='first')
df['_s'] = df['Supplier'].map(sup_rank)
df['_b'] = [brand_rank[(s, b)] for s, b in zip(df['Supplier'], df['brand'])]
df['_p'] = df['pnum'].map(prod_rank)
df = df.sort_values(['_s', '_b', '_p', 'lvl', 'start']).reset_index(drop=True)

N = len(df)
LAST = N + 1  # last data row on Deal Levels
tiers = sorted(df['lvl'].unique())
products = df.drop_duplicates('pnum')[['Supplier', 'pnum', 'pname', 'brand', 'pkg', 'family', '_s', '_b', '_p']].sort_values(['_s', '_b', '_p']).reset_index(drop=True)
NP = len(products)
PLAST = NP + 1
prod_tiers = df.groupby('pnum')['lvl'].apply(lambda s: sorted(set(s))).to_dict()
period = f"{df['start'].min():%m/%d/%Y} - {df['end'].max():%m/%d/%Y}"

# ----------------------------------------------------------------------------
# 2. Workbook helpers
# ----------------------------------------------------------------------------
wb = Workbook()

def hdr(ws, row, cols, fill=FILL_HDR, font=F_HDR):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font; c.fill = fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30

def put(ws, row, col, val, fmt=None, font=F_BODY, fill=None, bold=False):
    c = ws.cell(row=row, column=col, value=val)
    c.font = F_BOLD if bold else font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    return c

def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

# Deal Levels ranges (bounded, not whole-column, so LibreOffice stays quick)
DL = "'Deal Levels'"
def rng(col):
    return f"{DL}!${col}$2:${col}${LAST}"
R_SUP, R_PNUM, R_FAM, R_LVL, R_PRICE, R_CASES, R_GROSS, R_REV, R_BAND, R_BRAND, R_DISC = (
    rng('A'), rng('B'), rng('F'), rng('G'), rng('H'), rng('L'), rng('Q'), rng('S'), rng('X'), rng('D'), rng('V'))

# ----------------------------------------------------------------------------
# 3. README
# ----------------------------------------------------------------------------
ws = wb.active; ws.title = 'README'
lines = [
    ('Pricing / Deal Level Workbook -- Encompass "Pricing Analysis RDE by Month"', F_TITLE),
    (f'Source export: {os.path.basename(SRC)}   |   Deal windows in file: {period}   |   {N:,} export rows, {NP:,} products, {df["Supplier"].nunique()} suppliers', F_NOTE),
    ('', None),
    ('WHAT ONE ROW IS', F_BOLD),
    ('One export row = one product x one deal level (Min Purchase Qty, in cases) x one deal window, with the cases sold at that level in the window.', None),
    ('A product with 9 deal levels therefore appears on 9+ rows. Cases, Revenue and Gross sum cleanly across rows; prices and margins must be weighted by cases.', None),
    ('', None),
    ('DEFINITIONS (all from the export unless marked derived)', F_BOLD),
    ('Deal Level (Min Cases) = Min Purchase Qty: the quantity an account must buy to earn that Case Price. Level 0 = front-line / no minimum.', None),
    ('FOB = supplier invoice cost per case.  Tax = excise tax per case.  Participation = supplier-funded portion of the deal per case.', None),
    ('Laid-in Cost = the cost the margin is measured against. In 597 rows it does NOT equal FOB + Tax - Participation (flagged), so freight or other cost is embedded for some suppliers (Yuengling, Pabst, most spirits).', None),
    ('Gross $ = Cases x (Case Price - Laid-in Cost) as reported by Encompass. Margin % = Gross / (Cases x Case Price). Both are kept as reported: for spirits sold by the bottle the FOB/laid-in shown are PER BOTTLE while the price is per case, so recomputing from the columns would be wrong there.', None),
    ('Revenue $ (derived) = Cases x Case Price.  GP / Case (derived) = Gross / Cases.  Front-line Price (derived) = the highest Case Price the product carries in this export; Discount = Front-line - level price.', None),
    ('Level Band (derived) groups levels: 0 | 1-4 | 5-9 | 10-24 | 25-49 | 50-99 | 100-199 | 200+ cases.', None),
    ('Brand / Package (derived) are split off the product name at the first size token (e.g. "2/12/12 oz Btl", "15.5 Gal Keg"). Family (derived) is the inferred mix-and-match group -- see below.', None),
    ('', None),
    ('MIX & MATCH IS INFERRED, NOT IN THE EXPORT', F_BOLD),
    (f'The export has no deal-group / family field. Products were grouped within a supplier by the ladder of deal levels they share at {MM_MIN}+ cases (e.g. Corona bottles 25/54/108/162/320, Modelo cans 25/50/81/130/320, White Claw 25/50/91/104/208, Coors/Miller 25/35/50/150).', None),
    ('A product whose ladder matches no other product is folded into the supplier ladder it overlaps most (Jaccard >= 1/3) or left as its own family. Products with no 20+ case level are "no 20+ case level".', None),
    ('Treat the families as a strong hint of what mixes and matches, then CONFIRM against the supplier deal sheets / Encompass deal groups before using them in a pricing decision. Levels below 20 cases (2, 3, 5, 10) are more often single-SKU quantity breaks or account-type minimums.', None),
    ('', None),
    ('SHEETS', F_BOLD),
    ('Deal Levels        every export row, cleaned, with derived columns and data flags (filterable).', None),
    ('Price Ladder       one row per product: front-line vs deepest price, blended margin, and the CASE PRICE at every deal level (columns = level in cases).', None),
    ('Ladder - Cases     same grid, cases sold at each level.        Ladder - Margin %   same grid, margin % at each level.', None),
    ('M&M Families       one row per inferred family: ladder, products, cases, revenue, gross, margin.   M&M Members   which products sit in which family, with each product\'s levels.', None),
    ('Supplier Summary   cases / revenue / gross / margin by supplier, and how much volume moves at 25+ case levels.', None),
    ('Brand Summary      the same by brand.       Deal Depth   volume, margin and average discount by level band, overall and for the top suppliers.', None),
    ('Data Flags         rows worth a look: negative margins, no cost on file, laid-in that does not tie to FOB+Tax-Part, duplicate level/window rows (price overrides), credits.', None),
    ('', None),
    ('CAVEATS', F_BOLD),
    ('Aug-2026 only: one month cannot show seasonality, deal frequency or trend. 45 rows (0 cases) have no cost. 294 rows are negative-case credits and net against sales. Kegs carry no deposit in these numbers.', None),
    ('All totals on the summary sheets are live formulas over Deal Levels -- overwrite the data rows with a new export (same columns) and they recalculate; or rerun build_pricing_workbook.py.', None),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = f or F_BODY
    c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 160

# ----------------------------------------------------------------------------
# 4. Deal Levels
# ----------------------------------------------------------------------------
ws = wb.create_sheet('Deal Levels')
cols = ['Supplier', 'Product #', 'Product', 'Brand', 'Package', 'Family (inferred)', 'Deal Level (Min Cases)', 'Case Price',
        'Deal Start', 'Deal End', 'Days', 'Cases', 'FOB', 'Tax', 'Participation', 'Laid-in Cost', 'Gross $', 'Margin %',
        'Revenue $', 'GP / Case', 'Front-line Price', 'Discount $', 'Discount %', 'Level Band', 'Laid-in check (FOB+Tax-Part-Laid-in)', 'Flags']
hdr(ws, 1, cols)
PL = "'Price Ladder'"
for i, r in enumerate(df.itertuples(index=False), start=2):
    put(ws, i, 1, r.Supplier); put(ws, i, 2, int(r.pnum)); put(ws, i, 3, r.pname); put(ws, i, 4, r.brand); put(ws, i, 5, r.pkg)
    put(ws, i, 6, r.family); put(ws, i, 7, int(r.lvl), FMT_INT)
    put(ws, i, 8, float(r.price), FMT_PRICE)
    put(ws, i, 9, r.start.date(), FMT_DATE); put(ws, i, 10, r.end.date(), FMT_DATE)
    put(ws, i, 11, f'=J{i}-I{i}+1', FMT_INT)
    put(ws, i, 12, int(r.Cases), FMT_INT)
    put(ws, i, 13, None if pd.isna(r.FOB) else float(r.FOB), FMT_PRICE)
    put(ws, i, 14, float(r.Tax), FMT_PRICE)
    put(ws, i, 15, None if pd.isna(r.Participation) else float(r.Participation), FMT_PRICE)
    put(ws, i, 16, None if pd.isna(r.laidin) else float(r.laidin), FMT_PRICE)
    put(ws, i, 17, float(r.Gross), FMT_MONEY)
    put(ws, i, 18, float(r.margin), FMT_PCT)
    put(ws, i, 19, f'=H{i}*L{i}', FMT_MONEY)
    put(ws, i, 20, f'=IF(L{i}=0,"",Q{i}/L{i})', FMT_PRICE)
    put(ws, i, 21, f'=INDEX({PL}!$G$2:$G${PLAST},MATCH(B{i},{PL}!$B$2:$B${PLAST},0))', FMT_PRICE)
    put(ws, i, 22, f'=U{i}-H{i}', FMT_PRICE)
    put(ws, i, 23, f'=IF(U{i}=0,0,V{i}/U{i})', FMT_PCT)
    put(ws, i, 24, f'=IF(G{i}=0,"A. 0 (front-line)",IF(G{i}<5,"B. 1-4",IF(G{i}<10,"C. 5-9",IF(G{i}<25,"D. 10-24",IF(G{i}<50,"E. 25-49",IF(G{i}<100,"F. 50-99",IF(G{i}<200,"G. 100-199","H. 200+")))))))')
    put(ws, i, 25, f'=IF(P{i}="","",ROUND(M{i}+N{i}-O{i}-P{i},2))', FMT_PRICE)
    put(ws, i, 26, r.flags or None)
widths(ws, [26, 9, 44, 24, 22, 34, 10, 10, 10, 10, 6, 9, 9, 8, 10, 10, 12, 9, 12, 10, 10, 10, 9, 16, 14, 40])
ws.freeze_panes = 'D2'; ws.auto_filter.ref = f'A1:Z{LAST}'

# ----------------------------------------------------------------------------
# 5. Price Ladder + Cases + Margin grids
# ----------------------------------------------------------------------------
id_cols = ['Supplier', 'Product #', 'Product', 'Brand', 'Package', 'Family (inferred)']
sum_cols = ['Front-line Price (highest)', 'Deepest Price (lowest)', 'Max Discount $', 'Max Discount %', '# Levels',
            'Total Cases', 'Revenue $', 'Gross $', 'Blended Margin %', 'Cases at 25+ Levels', '% Cases at 25+']
FIRST_T = len(id_cols) + len(sum_cols) + 1  # first tier column index
LAST_T = FIRST_T + len(tiers) - 1
tcol = {t: FIRST_T + k for k, t in enumerate(tiers)}

def grid_sheet(title, grid_kind):
    ws = wb.create_sheet(title)
    hdr(ws, 1, id_cols + sum_cols + [f'{t}' for t in tiers])
    ws.cell(row=1, column=FIRST_T - 1).value = sum_cols[-1]
    for k, t in enumerate(tiers):
        ws.cell(row=1, column=FIRST_T + k, value=int(t)).number_format = '0'
    lab = {'price': 'CASE PRICE at each deal level (cases) -->', 'cases': 'CASES SOLD at each deal level (cases) -->', 'margin': 'MARGIN % at each deal level (cases) -->'}[grid_kind]
    for i, p in enumerate(products.itertuples(index=False), start=2):
        put(ws, i, 1, p.Supplier); put(ws, i, 2, int(p.pnum)); put(ws, i, 3, p.pname); put(ws, i, 4, p.brand); put(ws, i, 5, p.pkg); put(ws, i, 6, p.family)
        put(ws, i, 7, f'=_xlfn.MAXIFS({R_PRICE},{R_PNUM},B{i})', FMT_PRICE)
        put(ws, i, 8, f'=_xlfn.MINIFS({R_PRICE},{R_PNUM},B{i})', FMT_PRICE)
        put(ws, i, 9, f'=G{i}-H{i}', FMT_PRICE)
        put(ws, i, 10, f'=IF(G{i}=0,"",I{i}/G{i})', FMT_PCT)
        put(ws, i, 11, f'=COUNT({get_column_letter(FIRST_T)}{i}:{get_column_letter(LAST_T)}{i})' if grid_kind == 'price' else f"=INDEX('Price Ladder'!$K$2:$K${PLAST},MATCH(B{i},'Price Ladder'!$B$2:$B${PLAST},0))", FMT_INT)
        put(ws, i, 12, f'=SUMIFS({R_CASES},{R_PNUM},B{i})', FMT_INT)
        put(ws, i, 13, f'=SUMIFS({R_REV},{R_PNUM},B{i})', FMT_MONEY)
        put(ws, i, 14, f'=SUMIFS({R_GROSS},{R_PNUM},B{i})', FMT_MONEY)
        put(ws, i, 15, f'=IF(M{i}=0,"",N{i}/M{i})', FMT_PCT)
        put(ws, i, 16, f'=SUMIFS({R_CASES},{R_PNUM},B{i},{R_LVL},">=25")', FMT_INT)
        put(ws, i, 17, f'=IF(L{i}=0,"",P{i}/L{i})', FMT_PCT)
        for t in prod_tiers[p.pnum]:
            c = tcol[t]; L = get_column_letter(c)
            if grid_kind == 'price':
                put(ws, i, c, f'=AVERAGEIFS({R_PRICE},{R_PNUM},$B{i},{R_LVL},{L}$1)', FMT_PRICE)
            elif grid_kind == 'cases':
                put(ws, i, c, f'=SUMIFS({R_CASES},{R_PNUM},$B{i},{R_LVL},{L}$1)', FMT_INT)
            else:
                put(ws, i, c, f'=IFERROR(SUMIFS({R_GROSS},{R_PNUM},$B{i},{R_LVL},{L}$1)/SUMIFS({R_REV},{R_PNUM},$B{i},{R_LVL},{L}$1),AVERAGEIFS({rng("R")},{R_PNUM},$B{i},{R_LVL},{L}$1))', FMT_PCT)
    widths(ws, [24, 9, 42, 22, 20, 32, 10, 10, 9, 9, 7, 9, 11, 11, 9, 10, 9] + [7] * len(tiers))
    ws.freeze_panes = f'{get_column_letter(FIRST_T)}2'
    ws.auto_filter.ref = f'A1:{get_column_letter(LAST_T)}{PLAST}'
    note = ws.cell(row=PLAST + 2, column=1, value=f'{lab}  Column headers from "{get_column_letter(FIRST_T)}" on are the deal level in cases; a blank cell means the product had no rows at that level in this export. Front-line = highest price on any level.')
    note.font = F_NOTE
    return ws
grid_sheet('Price Ladder', 'price')
grid_sheet('Ladder - Cases', 'cases')
grid_sheet('Ladder - Margin %', 'margin')

# ----------------------------------------------------------------------------
# 6. M&M Families + Members
# ----------------------------------------------------------------------------
fam_df = products.groupby('family').agg(Supplier=('Supplier', 'first'), n=('pnum', 'size'), s=('_s', 'min')).reset_index()
fam_rev = df.groupby('family')['rev'].sum()
fam_df['rev'] = fam_df['family'].map(fam_rev)
fam_df = fam_df.sort_values(['s', 'rev'], ascending=[True, False]).reset_index(drop=True)
ws = wb.create_sheet('M&M Families')
hdr(ws, 1, ['Family (inferred)', 'Supplier', 'Ladder (deal levels >= 20 cases)', '# Products', 'Total Cases', 'Revenue $', 'Gross $', 'Blended Margin %', 'Cases at 25+ Levels', '% Cases at 25+', 'Cases at Front-line (level 0)', 'Avg Discount % (cases-wtd)', 'Products'])
for i, f in enumerate(fam_df.itertuples(index=False), start=2):
    ladder = f.family.split(' | ', 1)[1]
    plist = products[products['family'] == f.family]
    put(ws, i, 1, f.family); put(ws, i, 2, f.Supplier); put(ws, i, 3, ladder)
    put(ws, i, 4, f"=COUNTIFS('Price Ladder'!$F$2:$F${PLAST},A{i})", FMT_INT)
    put(ws, i, 5, f'=SUMIFS({R_CASES},{R_FAM},A{i})', FMT_INT)
    put(ws, i, 6, f'=SUMIFS({R_REV},{R_FAM},A{i})', FMT_MONEY)
    put(ws, i, 7, f'=SUMIFS({R_GROSS},{R_FAM},A{i})', FMT_MONEY)
    put(ws, i, 8, f'=IF(F{i}=0,"",G{i}/F{i})', FMT_PCT)
    put(ws, i, 9, f'=SUMIFS({R_CASES},{R_FAM},A{i},{R_LVL},">=25")', FMT_INT)
    put(ws, i, 10, f'=IF(E{i}=0,"",I{i}/E{i})', FMT_PCT)
    put(ws, i, 11, f'=SUMIFS({R_CASES},{R_FAM},A{i},{R_LVL},0)', FMT_INT)
    put(ws, i, 12, f'=IF(E{i}=0,"",SUMPRODUCT(({R_FAM}=A{i})*{R_CASES}*{rng("W")})/E{i})', FMT_PCT)
    put(ws, i, 13, '; '.join(f'{p} {n}' for p, n in zip(plist['pnum'], plist['pname']))[:32000])
widths(ws, [44, 26, 28, 9, 11, 12, 12, 10, 11, 9, 12, 11, 120])
ws.freeze_panes = 'B2'; ws.auto_filter.ref = f'A1:M{len(fam_df) + 1}'
r = len(fam_df) + 3
ws.cell(row=r, column=1, value=f'INFERRED from shared ladders of {MM_MIN}+ case levels within a supplier (see README). Confirm against supplier deal sheets. Avg Discount % is the cases-weighted discount off each product\'s front-line price.').font = F_NOTE

ws = wb.create_sheet('M&M Members')
hdr(ws, 1, ['Family (inferred)', 'Supplier', 'Product #', 'Product', 'Brand', 'Package', f'Levels >= {MM_MIN} cases', 'All Levels in Export', 'Total Cases', 'Revenue $', 'Gross $', 'Blended Margin %', 'Front-line Price', 'Deepest Price', 'Max Discount %'])
mem = products.merge(fam_df[['family']].reset_index().rename(columns={'index': '_f'}), on='family').sort_values(['_f', '_p'])
for i, p in enumerate(mem.itertuples(index=False), start=2):
    put(ws, i, 1, p.family); put(ws, i, 2, p.Supplier); put(ws, i, 3, int(p.pnum)); put(ws, i, 4, p.pname); put(ws, i, 5, p.brand); put(ws, i, 6, p.pkg)
    put(ws, i, 7, '/'.join(str(t) for t in prod_tiers[p.pnum] if t >= MM_MIN) or '-')
    put(ws, i, 8, '/'.join(str(t) for t in prod_tiers[p.pnum]))
    put(ws, i, 9, f'=SUMIFS({R_CASES},{R_PNUM},C{i})', FMT_INT)
    put(ws, i, 10, f'=SUMIFS({R_REV},{R_PNUM},C{i})', FMT_MONEY)
    put(ws, i, 11, f'=SUMIFS({R_GROSS},{R_PNUM},C{i})', FMT_MONEY)
    put(ws, i, 12, f'=IF(J{i}=0,"",K{i}/J{i})', FMT_PCT)
    put(ws, i, 13, f"=INDEX('Price Ladder'!$G$2:$G${PLAST},MATCH(C{i},'Price Ladder'!$B$2:$B${PLAST},0))", FMT_PRICE)
    put(ws, i, 14, f"=INDEX('Price Ladder'!$H$2:$H${PLAST},MATCH(C{i},'Price Ladder'!$B$2:$B${PLAST},0))", FMT_PRICE)
    put(ws, i, 15, f'=IF(M{i}=0,"",(M{i}-N{i})/M{i})', FMT_PCT)
widths(ws, [44, 26, 9, 44, 22, 20, 22, 30, 10, 12, 12, 10, 10, 10, 10])
ws.freeze_panes = 'D2'; ws.auto_filter.ref = f'A1:O{NP + 1}'

# ----------------------------------------------------------------------------
# 7. Supplier Summary / Brand Summary
# ----------------------------------------------------------------------------
def summary_sheet(title, key_label, keys, key_range, extra=None):
    ws = wb.create_sheet(title)
    cols = ([extra[0]] if extra else []) + [key_label, '# Products', '# Level Rows', 'Total Cases', 'Revenue $', 'Gross $', 'Margin %',
            'Cases at Front-line (0)', 'Cases at 1-24', 'Cases at 25+', '% Cases at 25+', 'Avg Discount % (cases-wtd)', 'GP / Case']
    hdr(ws, 1, cols)
    o = 1 if extra else 0
    K = get_column_letter(1 + o)
    for i, k in enumerate(keys, start=2):
        if extra: put(ws, i, 1, extra[1][k])
        put(ws, i, 1 + o, k)
        put(ws, i, 2 + o, f"=COUNTIFS('Price Ladder'!{extra[2] if extra else '$A$2:$A$' + str(PLAST)},{K}{i})", FMT_INT)
        put(ws, i, 3 + o, f'=COUNTIFS({key_range},{K}{i})', FMT_INT)
        put(ws, i, 4 + o, f'=SUMIFS({R_CASES},{key_range},{K}{i})', FMT_INT)
        put(ws, i, 5 + o, f'=SUMIFS({R_REV},{key_range},{K}{i})', FMT_MONEY)
        put(ws, i, 6 + o, f'=SUMIFS({R_GROSS},{key_range},{K}{i})', FMT_MONEY)
        c5, c6, c4 = get_column_letter(5 + o), get_column_letter(6 + o), get_column_letter(4 + o)
        put(ws, i, 7 + o, f'=IF({c5}{i}=0,"",{c6}{i}/{c5}{i})', FMT_PCT)
        put(ws, i, 8 + o, f'=SUMIFS({R_CASES},{key_range},{K}{i},{R_LVL},0)', FMT_INT)
        put(ws, i, 9 + o, f'=SUMIFS({R_CASES},{key_range},{K}{i},{R_LVL},">=1",{R_LVL},"<25")', FMT_INT)
        put(ws, i, 10 + o, f'=SUMIFS({R_CASES},{key_range},{K}{i},{R_LVL},">=25")', FMT_INT)
        c10 = get_column_letter(10 + o)
        put(ws, i, 11 + o, f'=IF({c4}{i}=0,"",{c10}{i}/{c4}{i})', FMT_PCT)
        put(ws, i, 12 + o, f'=IF({c4}{i}=0,"",SUMPRODUCT(({key_range}={K}{i})*{R_CASES}*{rng("W")})/{c4}{i})', FMT_PCT)
        put(ws, i, 13 + o, f'=IF({c4}{i}=0,"",{c6}{i}/{c4}{i})', FMT_PRICE)
    t = len(keys) + 2
    put(ws, t, 1 + o, 'TOTAL', bold=True, fill=FILL_TOT)
    if extra: put(ws, t, 1, '', fill=FILL_TOT)
    for c in range(2 + o, 14 + o):
        L = get_column_letter(c)
        if c in (7 + o, 11 + o, 12 + o, 13 + o):
            c5, c6, c4, c10 = get_column_letter(5 + o), get_column_letter(6 + o), get_column_letter(4 + o), get_column_letter(10 + o)
            f = {7 + o: f'={c6}{t}/{c5}{t}', 11 + o: f'={c10}{t}/{c4}{t}', 12 + o: f'=SUMPRODUCT({R_CASES}*{rng("W")})/{c4}{t}', 13 + o: f'={c6}{t}/{c4}{t}'}[c]
            put(ws, t, c, f, FMT_PCT if c != 13 + o else FMT_PRICE, bold=True, fill=FILL_TOT)
        else:
            put(ws, t, c, f'=SUM({L}2:{L}{t - 1})', FMT_MONEY if c in (5 + o, 6 + o) else FMT_INT, bold=True, fill=FILL_TOT)
    widths(ws, ([26] if extra else []) + [30, 9, 9, 11, 13, 13, 9, 11, 10, 10, 9, 11, 9])
    ws.freeze_panes = f'{get_column_letter(2 + o)}2'; ws.auto_filter.ref = f'A1:{get_column_letter(13 + o)}{t - 1}'
    return ws

sups = df.groupby('Supplier')['rev'].sum().sort_values(ascending=False).index.tolist()
summary_sheet('Supplier Summary', 'Supplier', sups, R_SUP)
brands = df.groupby('brand').agg(rev=('rev', 'sum'), sup=('Supplier', 'first')).sort_values('rev', ascending=False)
summary_sheet('Brand Summary', 'Brand', brands.index.tolist(), R_BRAND, extra=('Supplier', brands['sup'].to_dict(), f'$D$2:$D${PLAST}'))

# ----------------------------------------------------------------------------
# 8. Deal Depth (by level band)
# ----------------------------------------------------------------------------
ws = wb.create_sheet('Deal Depth')
bands = ['A. 0 (front-line)', 'B. 1-4', 'C. 5-9', 'D. 10-24', 'E. 25-49', 'F. 50-99', 'G. 100-199', 'H. 200+']
ws.cell(row=1, column=1, value='Volume, margin and discount by deal-level band -- all suppliers').font = F_TITLE
hdr(ws, 3, ['Level Band (cases)', '# Level Rows', 'Cases', '% of Cases', 'Revenue $', 'Gross $', 'Margin %', 'GP / Case', 'Avg Discount % (cases-wtd)', 'Avg Discount $ / Case'])
for i, b in enumerate(bands, start=4):
    put(ws, i, 1, b)
    put(ws, i, 2, f'=COUNTIFS({R_BAND},A{i})', FMT_INT)
    put(ws, i, 3, f'=SUMIFS({R_CASES},{R_BAND},A{i})', FMT_INT)
    put(ws, i, 4, f'=IF($C$12=0,"",C{i}/$C$12)', FMT_PCT)
    put(ws, i, 5, f'=SUMIFS({R_REV},{R_BAND},A{i})', FMT_MONEY)
    put(ws, i, 6, f'=SUMIFS({R_GROSS},{R_BAND},A{i})', FMT_MONEY)
    put(ws, i, 7, f'=IF(E{i}=0,"",F{i}/E{i})', FMT_PCT)
    put(ws, i, 8, f'=IF(C{i}=0,"",F{i}/C{i})', FMT_PRICE)
    put(ws, i, 9, f'=IF(C{i}=0,"",SUMPRODUCT(({R_BAND}=A{i})*{R_CASES}*{rng("W")})/C{i})', FMT_PCT)
    put(ws, i, 10, f'=IF(C{i}=0,"",SUMPRODUCT(({R_BAND}=A{i})*{R_CASES}*{R_DISC})/C{i})', FMT_PRICE)
t = 12
put(ws, t, 1, 'TOTAL', bold=True, fill=FILL_TOT)
for c, f, fm in [(2, '=SUM(B4:B11)', FMT_INT), (3, '=SUM(C4:C11)', FMT_INT), (4, '=SUM(D4:D11)', FMT_PCT), (5, '=SUM(E4:E11)', FMT_MONEY), (6, '=SUM(F4:F11)', FMT_MONEY),
                (7, '=F12/E12', FMT_PCT), (8, '=F12/C12', FMT_PRICE), (9, f'=SUMPRODUCT({R_CASES}*{rng("W")})/C12', FMT_PCT), (10, f'=SUMPRODUCT({R_CASES}*{R_DISC})/C12', FMT_PRICE)]:
    put(ws, t, c, f, fm, bold=True, fill=FILL_TOT)

TOPN = 15
top = sups[:TOPN]
r0 = 15
ws.cell(row=r0, column=1, value=f'Top {TOPN} suppliers x level band -- share of the supplier\'s cases').font = F_TITLE
hdr(ws, r0 + 2, ['Supplier'] + bands + ['Total Cases'])
for i, s in enumerate(top, start=r0 + 3):
    put(ws, i, 1, s)
    put(ws, i, 10, f'=SUMIFS({R_CASES},{R_SUP},A{i})', FMT_INT)
    for k, b in enumerate(bands, start=2):
        put(ws, i, k, f'=IF($J{i}=0,"",SUMIFS({R_CASES},{R_SUP},$A{i},{R_BAND},{get_column_letter(k)}${r0 + 2})/$J{i})', FMT_PCT)
r1 = r0 + 3 + TOPN + 2
ws.cell(row=r1, column=1, value=f'Top {TOPN} suppliers x level band -- margin % at each band').font = F_TITLE
hdr(ws, r1 + 2, ['Supplier'] + bands + ['Blended Margin %'])
for i, s in enumerate(top, start=r1 + 3):
    put(ws, i, 1, s)
    put(ws, i, 10, f'=IF(SUMIFS({R_REV},{R_SUP},A{i})=0,"",SUMIFS({R_GROSS},{R_SUP},A{i})/SUMIFS({R_REV},{R_SUP},A{i}))', FMT_PCT)
    for k, b in enumerate(bands, start=2):
        L = get_column_letter(k)
        put(ws, i, k, f'=IF(SUMIFS({R_REV},{R_SUP},$A{i},{R_BAND},{L}${r1 + 2})=0,"",SUMIFS({R_GROSS},{R_SUP},$A{i},{R_BAND},{L}${r1 + 2})/SUMIFS({R_REV},{R_SUP},$A{i},{R_BAND},{L}${r1 + 2}))', FMT_PCT)
r2 = r1 + 3 + TOPN + 1
ws.cell(row=r2, column=1, value='Bands come from the Level Band column on Deal Levels. Discount is off each product\'s front-line (highest) price in this export. Credits (negative cases) net against each band.').font = F_NOTE
widths(ws, [30, 14, 12, 12, 13, 13, 12, 12, 13, 12])
ws.freeze_panes = 'B4'

# ----------------------------------------------------------------------------
# 9. Data Flags
# ----------------------------------------------------------------------------
ws = wb.create_sheet('Data Flags')
fl = df[df['flags'] != ''].copy()
order = {'No cost on file': 0, 'Negative margin': 1, 'Duplicate level/window (price override?)': 2, 'Laid-in != FOB+Tax-Part': 3, 'Negative cases (credit / return)': 4}
fl['_o'] = fl['flags'].map(lambda s: min(order[x] for x in s.split('; ')))
fl = fl.sort_values(['_o', '_s', '_p', 'lvl'])
hdr(ws, 1, ['Flag(s)', 'Supplier', 'Product #', 'Product', 'Deal Level (Min Cases)', 'Case Price', 'Deal Start', 'Deal End', 'Cases', 'FOB', 'Tax', 'Participation', 'Laid-in Cost', 'FOB+Tax-Part', 'Laid-in delta', 'Gross $', 'Margin %'])
for i, r in enumerate(fl.itertuples(index=False), start=2):
    put(ws, i, 1, r.flags); put(ws, i, 2, r.Supplier); put(ws, i, 3, int(r.pnum)); put(ws, i, 4, r.pname)
    put(ws, i, 5, int(r.lvl), FMT_INT); put(ws, i, 6, float(r.price), FMT_PRICE)
    put(ws, i, 7, r.start.date(), FMT_DATE); put(ws, i, 8, r.end.date(), FMT_DATE)
    put(ws, i, 9, int(r.Cases), FMT_INT)
    put(ws, i, 10, None if pd.isna(r.FOB) else float(r.FOB), FMT_PRICE); put(ws, i, 11, float(r.Tax), FMT_PRICE)
    put(ws, i, 12, None if pd.isna(r.Participation) else float(r.Participation), FMT_PRICE)
    put(ws, i, 13, None if pd.isna(r.laidin) else float(r.laidin), FMT_PRICE)
    put(ws, i, 14, f'=IF(J{i}="","",J{i}+K{i}-L{i})', FMT_PRICE)
    put(ws, i, 15, f'=IF(N{i}="","",N{i}-M{i})', FMT_PRICE)
    put(ws, i, 16, float(r.Gross), FMT_MONEY); put(ws, i, 17, float(r.margin), FMT_PCT)
widths(ws, [44, 24, 9, 42, 10, 10, 10, 10, 8, 9, 8, 10, 10, 11, 10, 11, 9])
ws.freeze_panes = 'E2'; ws.auto_filter.ref = f'A1:Q{len(fl) + 1}'
cnt = Counter(x for s in df['flags'] if s for x in s.split('; '))
ws.cell(row=len(fl) + 3, column=1, value='Counts: ' + '; '.join(f'{k} = {v}' for k, v in cnt.items()) + '. Spirits sold by the bottle show per-bottle FOB / laid-in against a per-case price -- the reported Gross and Margin % are still right; the laid-in delta there is a units artefact, not an error.').font = F_NOTE

wb.save(OUT)
print(f'wrote {OUT}: {N} level rows, {NP} products, {len(fam_df)} inferred families, {len(fl)} flagged rows')
